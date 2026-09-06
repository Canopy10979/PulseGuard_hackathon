# Explain: Import pathlib import Path for document generation. Import docx import Document for document generation. Import docx.oxml.ns import qn for document generation.
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
# Explain: Import docx.shared import Inches, Pt, RGBColor for document generation. Assign ROOT from the value shown here. Assign OUT from the value shown here.
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentation" / "PulseGuard_Code_Presentation.docx"

# file, short purpose, exact source fragment, what that fragment proves
# Explain: Assign ITEMS from the value shown here. Record a presentation value or source-reference entry.
ITEMS = [
    ("index.html", "shows safety choices", "Open drop tracker", "the homepage exposes the tracker"),
    ("fall-detection.html", "holds tracker and care controls", "Find nearest care", "the page exposes the naloxones and hospital locator"),
    # Explain: Record a presentation value or source-reference entry.
    ("app.js", "runs detection and care lookup", "const [hospitals, naloxones] = await Promise.all([", "Google Places searches both resource types together"),
    ("auth.js", "controls accounts", 'const authDialog = document.querySelector("#auth-dialog");', "JavaScript selects the account dialog"),
    ("nav.js", "creates shared navigation", "const SITE = [", "all page links begin in one central structure"),
    # Explain: Record a presentation value or source-reference entry.
    ("sound.js", "handles ordinary interface audio", 'const SOUND_FILE = "click.mp3";', "click feedback stays separate from the warning alarm"),
    ("site.css", "defines shared visuals", "--bg: #101418;", "pages reuse one dark background token"),
    ("styles.css", "styles the homepage layout", ".choice-grid {", "the emergency choices use a dedicated grid"),
    # Explain: Record a presentation value or source-reference entry.
    ("about.html", "explains limits and privacy", "No breathing or pulse sensing.", "the site states a sensor limitation openly"),
    ("risk-zones.html", "demonstrates privacy-aware planning", "const MIN_CALLS = 5;", "small groups are hidden before a zone is drawn"),
    ("safewalk.html", "adds a timed check-in", "const GRACE_S = 300;", "late status includes a five-minute grace period"),
    # Explain: Record a presentation value or source-reference entry.
    ("shelters.html", "ranks nearby resources", "const NEAREST = 5;", "results are limited to five useful locations"),
    ("hazard.html", "merges nearby hazard reports", "const MERGE_M = 60;", "reports within sixty metres can become one issue"),
    ("script.js", "renders illustrative city data", "const CITIES = [", "the older dashboard reads from one defined dataset"),
# Explain: End this collection of values. Define the helper quote_for. Assign lines from the value shown here.
]

def quote_for(path, fragment):
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    # Explain: Repeat the following work for number, line in enumerate(lines, 1). Run the following block when fragment in line. Return number, line.strip() to the caller.
    for number, line in enumerate(lines, 1):
        if fragment in line:
            return number, line.strip()
    # Explain: Apply raise ValueError(f"Missing fragment in {path}: {fragment}") in the document-building sequence. Assign doc from the value shown here. Assign section from the value shown here.
    raise ValueError(f"Missing fragment in {path}: {fragment}")

doc = Document()
section = doc.sections[0]
# Explain: Assign section.page_width, section.page_height from the value shown here. Assign section.top_margin from the value shown here. Assign section.left_margin from the value shown here.
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.55)
section.left_margin = section.right_margin = Inches(.72)

# Explain: Assign styles from the value shown here. Set the font appearance for this text or style.
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
# Explain: Set the font appearance for this text or style. Set paragraph spacing or pagination behavior.
styles["Normal"].font.color.rgb = RGBColor(32, 36, 40)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.name = "Georgia"
# Explain: Set the font appearance for this text or style. Assign ppr from the value shown here.
styles["Title"].font.size = Pt(25)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
ppr = styles["Title"]._element.get_or_add_pPr()
# Explain: Repeat the following work for border in ppr.findall(qn("w:pBdr")). Apply ppr.remove(border) in the document-building sequence. Add the specified content or structure to the Word document.
for border in ppr.findall(qn("w:pBdr")):
    ppr.remove(border)

doc.add_paragraph("PulseGuard Code Presentation", style="Title")
# Explain: Add the specified content or structure to the Word document. Assign spoken from the value shown here. Repeat the following work for path, purpose, fragment, proof in ITEMS.
doc.add_paragraph("Approximately one minute thirty five seconds")

spoken = []
for path, purpose, fragment, proof in ITEMS:
    # Explain: Assign number, exact from the value shown here. Assign text from the value shown here. Apply spoken.append(text) in the document-building sequence.
    number, exact = quote_for(path, fragment)
    text = f"{path} {purpose}. Line {number}, \"{fragment}\", shows that {proof}."
    spoken.append(text)
    # Explain: Add the specified content or structure to the Word document. Assign p.add_run(path).bold from the value shown here. Apply p.add_run(text[len(path):]) in the document-building sequence.
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(path).bold = True
    p.add_run(text[len(path):])
    # Explain: Set paragraph spacing or pagination behavior. Assign word_count from the value shown here.
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(5)

word_count = sum(len(text.replace("—", " ").split()) for text in spoken)
# Explain: Add the specified content or structure to the Word document. Assign note.add_run("Timing ").bold from the value shown here. Apply note.add_run(f"{word_count} spoken words at about 170 words per minute.") in the document-building sequence.
note = doc.add_paragraph()
note.add_run("Timing  ").bold = True
note.add_run(f"{word_count} spoken words at about 170 words per minute.")

# Explain: Assign doc.core_properties.title from the value shown here. Assign doc.core_properties.subject from the value shown here. Assign doc.core_properties.author from the value shown here.
doc.core_properties.title = "PulseGuard Code Presentation"
doc.core_properties.subject = "Ninety five second code presentation with exact line references"
doc.core_properties.author = "PulseGuard Team"
# Explain: Write the completed Word document to its output path. Apply print(OUT) in the document-building sequence. Apply print(word_count) in the document-building sequence.
doc.save(OUT)
print(OUT)
print(word_count)
