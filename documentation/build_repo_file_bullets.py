from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentation" / "PulseGuard_Code_Presentation.docx"

# filename, feature, explanation, exact text used to locate the quoted line
ENTRIES = [
    (".gitignore", "keeps machine-specific files and secrets out of GitHub", "This protects repository hygiene without affecting the website", ".env"),
    ("README.md", "explains how to run and evaluate the project", "This gives another developer enough context to open the prototype correctly", "# PulseGuard"),
    ("about.html", "states the prototype's limits, privacy model, and evidence boundaries", "This keeps the safety claim focused on warning patterns instead of diagnosis", "<h1>What Lifeline can—and cannot—do.</h1>"),
    ("app.js", "controls the five-second arming delay, drop alarm, sensors, and nearby-care prompt", "This file contains the main safety sequence rather than only changing page appearance", "const ARM_DELAY_S = 5;"),
    ("auth.js", "opens the prototype login and sign-up dialog", "The code makes the account interface interactive without pretending that a backend exists", "if (!authDialog.open) authDialog.showModal();"),
    ("fall-detection.html", "provides the tracker button, alarm panel, Stop control, map prompt, and audio element", "This markup is the visible structure that app.js updates", '<audio id="alarmSound" src="warning-alarm.mp3" preload="auto" loop></audio>'),
    ("hazard.html", "demonstrates community hazard reporting and duplicate merging", "It is a supporting public-safety experiment rather than the core drop tracker", "const MERGE_M = 60;"),
    ("index.html", "presents the main emergency choices and a visible route to the tracker", "This makes the primary working feature easy to find from the homepage", '<a class="primary-btn" href="fall-detection.html">Open drop tracker</a>'),
    ("nav.js", "builds one shared navigation menu for every page", "Centralizing the link list prevents different pages from showing inconsistent destinations", 'const SITE = ['),
    ("risk-zones.html", "maps privacy-thresholded response-planning zones", "The page explores community planning while avoiding a claim about individual emergencies", "const MIN_CALLS = 5;"),
    ("robots.txt", "allows search crawlers to access the public site", "This is the smallest site-wide SEO instruction in the repository", "Allow: /"),
    ("safewalk.html", "demonstrates a timed check-in and nearby emergency-care search", "It supports the broader goal of shortening the route to human help", "const GRACE_S = 300;"),
    ("script.js", "renders the older illustrative city-data dashboard", "Its data stays clearly separate from the live drop-tracker logic", "const CITIES = ["),
    ("shelters.html", "ranks nearby AED, shelter, and hydrant records", "This supporting page shows how location could surface practical safety resources", "const NEAREST = 5;"),
    ("site.css", "defines shared navigation, controls, dialogs, accessibility states, and mobile rules", "This keeps the interface consistent across all pages", ".care-dialog {"),
    ("sound.js", "plays optional quiet feedback for ordinary interface actions", "It deliberately uses click.mp3 separately from the loud warning alarm", 'const SOUND_FILE = "click.mp3";'),
    ("styles.css", "styles homepage choices and responsive layout", "The grid becomes one column on smaller screens so the safety actions remain readable", ".choice-grid {"),
    ("documentation/build_prompt_aligned_presentation.py", "generates the prompt-relevance version of the Word presentation", "Keeping the generator in GitHub makes that document reproducible", 'OUT = r"'),
    ("documentation/build_pulseguard_105s_presentation.py", "generates the earlier 105-second bullet presentation", "Its talk list stores the exact presenter wording before Word formatting is applied", "talk = ["),
    ("documentation/build_two_minute_presentation.py", "generates the two-minute code overview", "It records the intended output path and can rebuild the document after code changes", 'OUT = r"'),
    ("documentation/build_repo_file_bullets.py", "generates this repository-wide line-reference presentation", "It reads the current files to calculate line numbers instead of relying on stale manual references", "ENTRIES = ["),
]

BINARY_ENTRIES = [
    ("click.mp3", "contains the short sound used for ordinary interface feedback", "As a binary audio file, it has no readable source-code line number"),
    ("warning-alarm.mp3", "contains the loud warning recording used after a detected drop", "As a binary audio file, it has no readable source-code line number"),
    ("documentation/Lifeline_Prompt_Aligned_Code_Presentation.docx", "stores the earlier prompt-aligned presentation", "A DOCX is a packaged binary document, so it has pages and paragraphs rather than source lines"),
    ("documentation/Lifeline_Two_Minute_Code_Presentation.docx", "stores the earlier two-minute presentation", "A DOCX is a packaged binary document, so it has pages and paragraphs rather than source lines"),
    ("documentation/PulseGuard_Code_Presentation.docx", "stores the current presentation deliverable", "A DOCX is a packaged binary document, so it has pages and paragraphs rather than source lines"),
]

def locate(path, needle):
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    for number, line in enumerate(lines, 1):
        if needle in line:
            return number, line.strip()
    raise ValueError(f"Could not find {needle!r} in {path}")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.55)
section.left_margin = section.right_margin = Inches(.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor(32, 36, 40)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.name = "Georgia"
styles["Title"].font.size = Pt(26)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
ppr = styles["Title"]._element.get_or_add_pPr()
for border in ppr.findall(qn("w:pBdr")):
    ppr.remove(border)
styles["Heading 1"].font.name = "Georgia"
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].font.bold = False
styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph("PulseGuard Repository File Presentation", style="Title")
doc.add_paragraph("Two to three presentation sentences for every file tracked in GitHub")
intro = doc.add_paragraph()
intro.add_run("How to use this document  ").bold = True
intro.add_run("Each bullet identifies a file's feature, quotes an exact current line when the file is readable text, and explains what that line proves about the implementation.")

doc.add_heading("Source and configuration files", level=1)
for path, feature, explanation, needle in ENTRIES:
    number, exact = locate(path, needle)
    p = doc.add_paragraph(style="List Bullet")
    first = p.add_run(path + ". ")
    first.bold = True
    p.add_run(f"This file {feature}. Line {number}, \"{exact}\", shows the exact code used for that feature. {explanation}.")
    p.paragraph_format.keep_together = True

doc.add_heading("Binary files", level=1)
for path, feature, explanation in BINARY_ENTRIES:
    p = doc.add_paragraph(style="List Bullet")
    first = p.add_run(path + ". ")
    first.bold = True
    p.add_run(f"This file {feature}. {explanation}.")
    p.paragraph_format.keep_together = True

doc.core_properties.title = "PulseGuard Repository File Presentation"
doc.core_properties.subject = "Exact code-line explanations for every file in the GitHub repository"
doc.core_properties.author = "PulseGuard Team"
doc.save(OUT)
print(OUT)
